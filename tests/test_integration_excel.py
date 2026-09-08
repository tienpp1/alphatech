"""
Unit & Integration Tests for Excel (.xlsx) Parser & Ingestion.
"""

import io
from datetime import datetime, date
import openpyxl
from django.test import TestCase
from apps.integration.parsers.excel_parser import parse_excel_file


def _create_sample_excel_bytes(sheet_name="Data", rows=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    if rows:
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


class ExcelParserTests(TestCase):
    def test_valid_excel_parsing(self):
        rows = [
            ["ma_kh", "ho_ten", "so_tien", "ngay_tao", "trang_thai"],
            ["KH-01", "Nguyen Van A", 1500000, datetime(2026, 8, 20, 10, 0), True],
            ["KH-02", "Tran Thi B", 2200000, date(2026, 8, 21), False],
        ]
        excel_bytes = _create_sample_excel_bytes("Orders", rows)
        res = parse_excel_file(excel_bytes)

        self.assertTrue(res["is_valid"])
        self.assertEqual(res["total_rows"], 2)
        self.assertEqual(res["columns"], ["ma_kh", "ho_ten", "so_tien", "ngay_tao", "trang_thai"])
        self.assertEqual(res["rows"][0]["raw_data"]["ma_kh"], "KH-01")
        self.assertEqual(res["rows"][0]["raw_data"]["so_tien"], 1500000)
        self.assertEqual(res["rows"][0]["raw_data"]["trang_thai"], True)

    def test_excel_multi_sheet_selection(self):
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Summary"
        ws1.append(["title", "value"])
        ws1.append(["Report", 123])

        ws2 = wb.create_sheet(title="Customers")
        ws2.append(["customer_id", "full_name"])
        ws2.append(["C-001", "Alice"])
        ws2.append(["C-002", "Bob"])

        buf = io.BytesIO()
        wb.save(buf)
        excel_bytes = buf.getvalue()

        # Parse specific sheet
        res = parse_excel_file(excel_bytes, sheet_name="Customers")
        self.assertTrue(res["is_valid"])
        self.assertEqual(res["sheet_name"], "Customers")
        self.assertEqual(res["total_rows"], 2)
        self.assertEqual(res["columns"], ["customer_id", "full_name"])

    def test_excel_invalid_sheet_name_error(self):
        excel_bytes = _create_sample_excel_bytes("SheetA", [["h1", "h2"], ["v1", "v2"]])
        res = parse_excel_file(excel_bytes, sheet_name="NonExistentSheet")

        self.assertFalse(res["is_valid"])
        self.assertIn("does not exist", res["error_message"])

    def test_empty_excel_file(self):
        wb = openpyxl.Workbook()
        buf = io.BytesIO()
        wb.save(buf)
        excel_bytes = buf.getvalue()

        res = parse_excel_file(excel_bytes)
        self.assertFalse(res["is_valid"])
        self.assertEqual(res["total_rows"], 0)

    def test_corrupted_excel_bytes(self):
        res = parse_excel_file(b"This is not a valid zip or xlsx file content.")
        self.assertFalse(res["is_valid"])
        self.assertIn("Invalid or corrupted Excel file", res["error_message"])

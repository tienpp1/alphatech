"""Create an academically precise copy of the teacher-reviewed proposal.

The source DOCX is never overwritten.  This script only corrects claims that
would otherwise overstate the implementation and normalizes the online IEEE
references for the final review copy.
"""

from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


SOURCE = Path(
    r"D:\ai_business_platform\output\docx_review\teacher_review_v2\De_cuong_Ha_Minh_Tien_sua_gop_y_lan_2.docx"
)
TARGET = Path(
    r"D:\ai_business_platform\output\docx_review\teacher_review_v2\De_cuong_Ha_Minh_Tien_doi_chieu_hoc_thuat.docx"
)


REPLACEMENTS = {
    "chat realtime, bản tin nội bộ": (
        "chat gần thời gian thực bằng polling 3 giây, bản tin nội bộ"
    ),
    "Chat realtime hỗ trợ nhân viên và người quản lý trao đổi trực tiếp": (
        "Chat gần thời gian thực hỗ trợ nhân viên và người quản lý trao đổi "
        "bằng cơ chế polling định kỳ 3 giây"
    ),
    "tập huấn luyện, kiểm định và kiểm thử": (
        "tập huấn luyện và kiểm thử theo thứ tự thời gian; chỉ ghi tập kiểm "
        "định riêng khi cấu hình thực nghiệm thực sự tạo tập này"
    ),
    "Đối với chat realtime, kiểm thử": (
        "Đối với chat gần thời gian thực bằng polling 3 giây, kiểm thử"
    ),
    "trao đổi tin nhắn realtime giữa": (
        "trao đổi tin nhắn gần thời gian thực bằng polling 3 giây giữa"
    ),
    "chat realtime và bản tin nội bộ": (
        "chat gần thời gian thực bằng polling 3 giây và bản tin nội bộ"
    ),
}


ONLINE_PATTERN = re.compile(
    r"\s*\[Trực tuyến\]\. Địa chỉ: (?P<url>https?://\S+?)\. "
    r"\[Truy cập: \d{2}/\d{2}/\d{4}\]\."
)


def replace_paragraph_text(paragraph, text):
    """Replace a paragraph while retaining the first run's character style."""
    if paragraph.runs:
        first = paragraph.runs[0]
        first.text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def iter_all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    document = Document(SOURCE)
    replacement_count = 0
    online_reference_count = 0

    for paragraph in iter_all_paragraphs(document):
        updated = paragraph.text
        for original, replacement in REPLACEMENTS.items():
            if original in updated:
                updated = updated.replace(original, replacement)
                replacement_count += 1

        if paragraph.text.lstrip().startswith("["):
            match = ONLINE_PATTERN.search(updated)
            if match:
                url = match.group("url")
                updated = ONLINE_PATTERN.sub(
                    f" [Online]. Available: {url}. [Accessed: Sep. 26, 2026].",
                    updated,
                )
                online_reference_count += 1

            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.left_indent = Cm(1.0)
            paragraph.paragraph_format.first_line_indent = Cm(-1.0)
            paragraph.paragraph_format.space_after = Pt(6)

        if updated != paragraph.text:
            replace_paragraph_text(paragraph, updated)

    if replacement_count != 7:
        raise RuntimeError(
            f"Expected 7 implementation-accuracy replacements, got {replacement_count}."
        )
    if online_reference_count != 10:
        raise RuntimeError(
            f"Expected 10 online IEEE references, got {online_reference_count}."
        )

    document.core_properties.title = (
        "Đề cương đồ án chuyên ngành — bản đối chiếu học thuật"
    )
    document.core_properties.comments = (
        "Bản sao từ teacher_review_v2; làm rõ polling 3 giây, chia tập dự báo "
        "và chuẩn hóa tài liệu trực tuyến IEEE ngày 26/09/2026."
    )
    TARGET.parent.mkdir(parents=True, exist_ok=True)
    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
